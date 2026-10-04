import io

import pytest
from docx import Document
from fpdf import FPDF

from app.attachments.extractor import (
    AttachmentExtractionError,
    UnsupportedAttachmentError,
    enrich_transcript,
    extract_text,
)


def make_pdf_bytes(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    return bytes(pdf.output())


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_text_reads_pdf_content():
    content = make_pdf_bytes("Budget approved: 50000 EUR")

    text = extract_text(filename="budget.pdf", content=content, max_chars=10_000)

    assert "Budget approved: 50000 EUR" in text


def test_extract_text_reads_docx_content_and_skips_empty_paragraphs():
    content = make_docx_bytes(["Project kickoff notes.", "", "Team: 3 developers."])

    text = extract_text(filename="notes.docx", content=content, max_chars=10_000)

    assert "Project kickoff notes." in text
    assert "Team: 3 developers." in text


def test_extract_text_truncates_to_max_chars():
    content = make_docx_bytes(["x" * 100])

    text = extract_text(filename="long.docx", content=content, max_chars=10)

    assert len(text) == 10


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(UnsupportedAttachmentError):
        extract_text(filename="notes.txt", content=b"hello world", max_chars=1000)


def test_extract_text_raises_extraction_error_on_corrupt_pdf():
    with pytest.raises(AttachmentExtractionError):
        extract_text(filename="broken.pdf", content=b"not a pdf", max_chars=1000)


def test_enrich_transcript_delimits_each_attachment():
    result = enrich_transcript(
        transcript="Kickoff call transcript.",
        attachments=[("budget.pdf", "50000 EUR total")],
    )

    assert "Kickoff call transcript." in result
    assert "--- attachment: budget.pdf ---" in result
    assert "50000 EUR total" in result
    assert "--- end attachment ---" in result


def test_enrich_transcript_without_attachments_returns_transcript_unchanged():
    result = enrich_transcript(transcript="Just a transcript.", attachments=[])

    assert result == "Just a transcript."
